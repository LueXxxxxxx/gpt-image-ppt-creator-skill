#!/usr/bin/env python3
"""Package full-slide images into an image-only PowerPoint deck."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches
except ModuleNotFoundError as exc:
    raise SystemExit(
        "Missing dependency: python-pptx. Install it with "
        "`python -m pip install -r scripts/requirements.txt`."
    ) from exc


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


def natural_key(path: Path) -> list[object]:
    parts = re.split(r"(\d+)", path.name.lower())
    return [int(part) if part.isdigit() else part for part in parts]


def collect_images(images_dir: Path) -> list[Path]:
    images = [
        path
        for path in images_dir.iterdir()
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
    ]
    return sorted(images, key=natural_key)


def build_pptx(images: list[Path], output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    if len(prs.slides) == 1:
        xml_slides = prs.slides._sldIdLst  # noqa: SLF001
        rel_id = xml_slides[0].rId
        prs.part.drop_rel(rel_id)
        xml_slides.remove(xml_slides[0])

    for image in images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(image),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--images", required=True, type=Path, help="Directory of slide images")
    parser.add_argument("--output", required=True, type=Path, help="Output .pptx path")
    args = parser.parse_args()

    if not args.images.exists() or not args.images.is_dir():
        parser.error(f"Image directory does not exist: {args.images}")

    images = collect_images(args.images)
    if not images:
        parser.error(f"No supported images found in: {args.images}")

    build_pptx(images, args.output)
    print(f"Created {args.output} with {len(images)} slides")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
